# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

prs = Presentation('document/Bao_Cao_Do_An_RDM6300_PicoRV32_SoC.pptx')
print(f"Total Slides: {len(prs.slides)}")

for s_num in [8, 9]:
    s = prs.slides[s_num - 1]
    print(f"\n=== SLIDE {s_num:02d} ===")
    for shp in s.shapes:
        if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
            print(f"  [PICTURE] Left={shp.left.inches:.2f}\", Top={shp.top.inches:.2f}\", Width={shp.width.inches:.2f}\", Height={shp.height.inches:.2f}\"")
        elif shp.has_text_frame:
            for p in shp.text_frame.paragraphs:
                txt = p.text.strip()
                if txt and ("CHƯƠNG" in txt.upper() or "CHU TRÌNH" in txt.upper() or "LUỒNG" in txt.upper() or "BƯỚC" in txt.upper()):
                    print(f"  [TEXT] {txt[:90]}")
