import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

out_dir = r"c:\Users\HP DRAGONFLY G2\Desktop\uart-fifo-rfid-asic\rdm6300-picorv32-rom-data\document\temp"
logo_dir = os.path.join(out_dir, "logos")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

C_NAVY_DARK   = RGBColor(11, 19, 43)
C_NAVY_MID    = RGBColor(28, 37, 65)
C_BLUE_ACCENT = RGBColor(37, 99, 235)
C_CYAN_ACCENT = RGBColor(14, 165, 233)
C_WHITE       = RGBColor(255, 255, 255)

s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = C_NAVY_DARK
bg1.line.fill.background()

card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
card1.fill.solid()
card1.fill.fore_color.rgb = C_NAVY_MID
card1.line.color.rgb = C_BLUE_ACCENT
card1.line.width = Pt(2.0)

tb1 = s1.shapes.add_textbox(Inches(1.05), Inches(1.05), Inches(11.233), Inches(4.2))
tf1 = tb1.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_right = tf1.margin_top = tf1.margin_bottom = 0

p_org = tf1.paragraphs[0]
p_org.text = "FPT JETKING — CHIP DESIGN"
p_org.font.name = "Segoe UI"
p_org.font.size = Pt(13.5)
p_org.font.bold = True
p_org.font.color.rgb = C_CYAN_ACCENT
p_org.space_after = Pt(12)

p_main = tf1.add_paragraph()
p_main.text = "THIẾT KẾ HỆ THỐNG SOC XỬ LÝ DỮ LIỆU THẺ RA VÀO RFID\nTỐI ƯU RTL TO GDSII TRÊN CHIP ASIC"
p_main.font.name = "Segoe UI"
p_main.font.size = Pt(26.0)
p_main.font.bold = True
p_main.font.color.rgb = C_WHITE
p_main.space_after = Pt(14)
p_main.line_spacing = 1.25

p_sub = tf1.add_paragraph()
p_sub.text = "Giao tiếp RFID RDM6300 | Lưu trữ Whitelist SPI Flash NVM | Tạo mẫu Basys 3 FPGA | Ký duyệt ASIC OpenLane 2 (Sky130)"
p_sub.font.name = "Segoe UI"
p_sub.font.size = Pt(12.0)
p_sub.font.color.rgb = RGBColor(203, 213, 225)
p_sub.space_after = Pt(18)

p_info1 = tf1.add_paragraph()
p_info1.text = "Học viên thực hiện :  Thái Tuấn Hiệp"
p_info1.font.name = "Segoe UI"
p_info1.font.size = Pt(13.0)
p_info1.font.bold = True
p_info1.font.color.rgb = C_WHITE
p_info1.space_after = Pt(5)

p_info2 = tf1.add_paragraph()
p_info2.text = "Giảng viên hướng dẫn :  ThS. Nguyễn Văn Đông"
p_info2.font.name = "Segoe UI"
p_info2.font.size = Pt(12.5)
p_info2.font.color.rgb = RGBColor(226, 232, 240)
p_info2.space_after = Pt(5)

p_info3 = tf1.add_paragraph()
p_info3.text = "Học kỳ :  SEM3  |  Chuyên ngành: Thiết kế Vi mạch Bán dẫn (Chip Design)"
p_info3.font.name = "Segoe UI"
p_info3.font.size = Pt(11.5)
p_info3.font.italic = True
p_info3.font.color.rgb = RGBColor(148, 163, 184)

# Logo header label
tb_lbl = s1.shapes.add_textbox(Inches(1.05), Inches(5.35), Inches(11.233), Inches(0.28))
tf_lbl = tb_lbl.text_frame
tf_lbl.margin_left = tf_lbl.margin_right = tf_lbl.margin_top = tf_lbl.margin_bottom = 0
p_l = tf_lbl.paragraphs[0]
p_l.text = "CÔNG CỤ EDA & HỆ SINH THÁI CÔNG NGHỆ VI MẠCH SỬ DỤNG TRONG ĐỒ ÁN:"
p_l.font.name = "Segoe UI"
p_l.font.size = Pt(10.5)
p_l.font.bold = True
p_l.font.color.rgb = C_CYAN_ACCENT

# 5 Logo badges
badges = [
    ("badge_openlane.png", 247, 87),
    ("badge_vivado.png", 240, 87),
    ("badge_fpga.png", 254, 87),
    ("badge_riscv.png", 381, 87),
    ("badge_skywater.png", 290, 87),
]

b_h = Inches(0.72)
badge_widths = [b_h * (w / h) for _, w, h in badges]
total_bw = sum(badge_widths)
avail_w = Inches(11.233)
gap = (avail_w - total_bw) / (len(badges) - 1)
cur_x = Inches(1.05)
badge_y = Inches(5.72)

for (b_file, _, _), bw in zip(badges, badge_widths):
    b_path = os.path.join(logo_dir, b_file)
    if os.path.exists(b_path):
        s1.shapes.add_picture(b_path, cur_x, badge_y, bw, b_h)
    cur_x += bw + gap

test_pptx = os.path.join(out_dir, "test_slide1_logos.pptx")
prs.save(test_pptx)
print("Saved test slide 1 pptx!")
