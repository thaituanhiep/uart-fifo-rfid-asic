# -*- coding: utf-8 -*-
"""
Script: build_13slides.py
Tạo bộ slide báo cáo PowerPoint 16:9 widescreen 13 slide tinh gọn:
- Giữ lại duy nhất Slide Sơ đồ khối (Block Diagram) cho phần kiến trúc SoC.
- Trên slide Sơ đồ khối, chia 3 bước trực quan:
  + Bước 1: Các khối nền tảng có sẵn (picorv32.v, spimemio.v, data_sram)
  + Bước 2: Cấu hình hệ thống chạy firmware (soc_interconnect, sections.lds, start.s)
  + Bước 3: Lập trình C ứng dụng MMIO (soc_regs.h, access_control.c)
- Lược bỏ hoàn toàn các slide code và bảng chu kỳ bus phức tạp.
- Tổng cộng: Đúng 13 slide chuẩn mực, chuyên nghiệp, không gây áp lực khi thuyết trình.
"""

import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Bảng màu thiết kế chuyên nghiệp
    C_NAVY_DARK   = RGBColor(11, 19, 43)     # #0B132B
    C_NAVY_MID    = RGBColor(28, 37, 65)     # #1C2541
    C_BLUE_ACCENT = RGBColor(37, 99, 235)    # #2563EB
    C_CYAN_ACCENT = RGBColor(14, 165, 233)   # #0EA5E9
    C_BG_LIGHT    = RGBColor(248, 250, 252)  # #F8FAFC
    C_WHITE       = RGBColor(255, 255, 255)  # #FFFFFF
    C_CARD_BG     = RGBColor(255, 255, 255)  # #FFFFFF
    C_CARD_BORDER = RGBColor(203, 213, 225)  # #CBD5E1
    C_TEXT_DARK   = RGBColor(15, 23, 42)     # #0F172A
    C_TEXT_MUTED  = RGBColor(71, 85, 105)    # #475569
    C_GREEN       = RGBColor(16, 185, 129)   # #10B981
    C_AMBER       = RGBColor(217, 119, 6)    # #D97706
    C_PURPLE      = RGBColor(124, 58, 237)   # #7C3AED
    C_ROSE        = RGBColor(225, 29, 72)    # #E11D48

    cur_dir = os.path.dirname(os.path.abspath(__file__))
    doc_dir = os.path.dirname(cur_dir) if os.path.basename(cur_dir) == "temp" else cur_dir
    project_root = os.path.dirname(doc_dir)

    img_fig1 = os.path.join(cur_dir, "fig1_block_diagram.png")
    if not os.path.exists(img_fig1):
        img_fig1 = os.path.join(doc_dir, "fig1_block_diagram.png")

    img_openroad = os.path.join(project_root, "OpenROAD.png")
    if not os.path.exists(img_openroad):
        img_openroad = os.path.join(cur_dir, "OpenROAD.png")
    if not os.path.exists(img_openroad):
        img_openroad = os.path.join(doc_dir, "OpenROAD.png")

    img_signoff = os.path.join(project_root, "AntennaLvsDrc.png")
    if not os.path.exists(img_signoff):
        img_signoff = os.path.join(cur_dir, "AntennaLvsDrc.png")
    if not os.path.exists(img_signoff):
        img_signoff = os.path.join(doc_dir, "AntennaLvsDrc.png")

    TOTAL_SLIDES = 13

    # Helper: Tiêu đề slide chuẩn hóa
    def add_header(slide, category, title, slide_num, total_slides=TOTAL_SLIDES):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_LIGHT
        bg.line.fill.background()

        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_BLUE_ACCENT
        top_bar.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.22), Inches(11.733), Inches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_BLUE_ACCENT
        p_cat.space_after = Pt(2)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(19)
        p_title.font.bold = True
        p_title.font.color.rgb = C_TEXT_DARK

        footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.02))
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = RGBColor(226, 232, 240)
        footer_line.line.fill.background()

        tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(8.0), Inches(0.3))
        tf_foot = tb_foot.text_frame
        tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "Đồ Án Tốt Nghiệp: SoC PicoRV32 RFID RDM6300 & SPI Flash | SV: Thái Tuấn Hiệp | GVHD: ThS. Nguyễn Văn Đông"
        p_foot.font.name = "Segoe UI"
        p_foot.font.size = Pt(9.5)
        p_foot.font.color.rgb = C_TEXT_MUTED

        tb_num = slide.shapes.add_textbox(Inches(11.0), Inches(7.1), Inches(1.533), Inches(0.3))
        tf_num = tb_num.text_frame
        tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.text = f"{slide_num:02d} / {total_slides:02d}"
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.name = "Segoe UI"
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = C_BLUE_ACCENT

    # Helper: Hộp nội dung có viền (Card)
    def add_card(slide, x, y, w, h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=1.0):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
        return card

    # Helper: Hộp mã nguồn dạng Terminal
    def add_code_box(slide, x, y, w, h, title, lines, status_text=None):
        card = add_card(slide, x, y, w, h, bg_color=C_NAVY_DARK, border_color=C_NAVY_MID, border_width=1.5)
        tb_t = slide.shapes.add_textbox(x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), Inches(0.35))
        tf_t = tb_t.text_frame
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = f">_  {title}"
        p_t.font.name = "Consolas"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = C_CYAN_ACCENT

        h_code = h - Inches(0.5) if not status_text else h - Inches(0.85)
        tb_c = slide.shapes.add_textbox(x + Inches(0.15), y + Inches(0.45), w - Inches(0.3), h_code)
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

        for idx, line in enumerate(lines):
            p = tf_c.paragraphs[0] if idx == 0 else tf_c.add_paragraph()
            p.text = line
            p.font.name = "Consolas"
            p.font.size = Pt(8.0)
            if "[PASS]" in line or "100%" in line or "CONFIRMED" in line or "MET TIMING" in line:
                p.font.color.rgb = C_GREEN
                p.font.bold = True
            elif "[TEST" in line or "[BOOT]" in line or "[EXEC]" in line or "[UART]" in line or "[HOST]" in line:
                p.font.color.rgb = C_CYAN_ACCENT
                p.font.bold = True
            elif "Analyzing" in line or "//" in line:
                p.font.color.rgb = RGBColor(148, 163, 184)
            else:
                p.font.color.rgb = C_WHITE

        if status_text:
            tb_s = slide.shapes.add_textbox(x + Inches(0.15), y + h - Inches(0.4), w - Inches(0.3), Inches(0.3))
            tf_s = tb_s.text_frame
            tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
            p_s = tf_s.paragraphs[0]
            p_s.text = status_text
            p_s.font.name = "Consolas"
            p_s.font.size = Pt(9.5)
            p_s.font.bold = True
            p_s.font.color.rgb = C_GREEN

    # =========================================================================
    # SLIDE 1: COVER / TITLE SLIDE (Dark Tech Widescreen)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = C_NAVY_DARK
    bg1.line.fill.background()

    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_NAVY_MID
    card1.line.color.rgb = C_BLUE_ACCENT
    card1.line.width = Pt(2.0)

    tb1 = s1.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(10.9), Inches(5.3))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "BÁO CÁO ĐỒ ÁN THIẾT KẾ VI MẠCH SỐ & HỆ THỐNG NHÚNG"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = C_CYAN_ACCENT
    p0.space_after = Pt(10)

    p_main = tf1.add_paragraph()
    p_main.text = "THIẾT KẾ HỆ THỐNG TRÊN VI MẠCH (SoC) CHO THIẾT BỊ KIỂM SOÁT RA VÀO DỰA TRÊN RFID RDM6300 VÀ BỘ NHỚ SPI FLASH"
    p_main.font.name = "Segoe UI"
    p_main.font.size = Pt(22)
    p_main.font.bold = True
    p_main.font.color.rgb = C_WHITE
    p_main.space_after = Pt(14)
    p_main.line_spacing = 1.15

    p_sub = tf1.add_paragraph()
    p_sub.text = "Tích hợp CPU RISC-V PicoRV32 | Kiến trúc Boot Flash XIP | Tạo mẫu Basys 3 FPGA | Ký duyệt ASIC OpenLane 2 (SkyWater 130nm)"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(12)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)
    p_sub.space_after = Pt(24)

    p_info1 = tf1.add_paragraph()
    p_info1.text = "Sinh viên thực hiện :  Thái Tuấn Hiệp  -  MSSV: 20210328  (Chuyên ngành Kỹ thuật Vi điện tử)"
    p_info1.font.name = "Segoe UI"
    p_info1.font.size = Pt(12.5)
    p_info1.font.bold = True
    p_info1.font.color.rgb = C_WHITE
    p_info1.space_after = Pt(6)

    p_info2 = tf1.add_paragraph()
    p_info2.text = "Giảng viên hướng dẫn :  ThS. Nguyễn Văn Đông  -  Bộ môn Kỹ thuật Máy tính & Vi điện tử"
    p_info2.font.name = "Segoe UI"
    p_info2.font.size = Pt(12)
    p_info2.font.color.rgb = RGBColor(226, 232, 240)
    p_info2.space_after = Pt(6)

    p_info3 = tf1.add_paragraph()
    p_info3.text = "Thời gian thực hiện :  Học kỳ 2025.2 - 2026.1  |  Địa điểm: Phòng thí nghiệm Thiết kế Vi mạch VLSI"
    p_info3.font.name = "Segoe UI"
    p_info3.font.size = Pt(11)
    p_info3.font.italic = True
    p_info3.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: MỤC LỤC / NỘI DUNG BÁO CÁO (Agenda - 6 Phần Cốt Lõi)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Tổng Quan Báo Cáo", "Nội Dung Báo Cáo Đồ Án (6 Phần Chuẩn Mực)", 2, total_slides=TOTAL_SLIDES)

    agenda_items = [
        ("Phần 1: Giới Thiệu Dự Án", "Hệ thống kiểm soát ra vào Offline bảo mật cao, lý do chọn bộ nhớ SPI Flash NVM."),
        ("Phần 2: Nền Tảng Phần Cứng & Phần Mềm", "Kit FPGA Basys 3, đầu đọc RDM6300, chuỗi công cụ Vivado, OpenLane 2 & RISC-V GCC."),
        ("Phần 3: Kiến Trúc Vi Hệ Thống SoC", "Sơ đồ khối Block Diagram chuẩn xuất bản và 3 bước hiện thực Co-Design phần cứng - firmware."),
        ("Phần 4: Hệ Thống Testbench & Mô Phỏng", "Kiểm thử RTL thuần (tb_uart_rtl.v) và đồng mô phỏng toàn diện Top SoC Boot Flash (tb_uart_ping.v)."),
        ("Phần 5: Demo Thực Nghiệm Trên FPGA", "Tạo mẫu thực tế trên kit Basys 3, giao tiếp UART PC và 4 kịch bản quẹt thẻ RFID thật."),
        ("Phần 6: Thiết Kế Vật Lý ASIC & Ký Duyệt", "Quy trình RTL-to-GDSII trên OpenLane 2 (Sky130): Layout GDSII, chỉ số PPA và ký duyệt Sign-off.")
    ]

    for i, (ag_t, ag_d) in enumerate(agenda_items):
        col = i % 3
        row = i // 3
        ax = Inches(0.8) + col * Inches(3.95)
        ay = Inches(1.50) + row * Inches(2.65)
        aw = Inches(3.75)
        ah = Inches(2.35)

        card_ag = add_card(s2, ax, ay, aw, ah, bg_color=C_WHITE, border_color=C_BLUE_ACCENT, border_width=1.5)
        tb_ag = s2.shapes.add_textbox(ax + Inches(0.20), ay + Inches(0.20), aw - Inches(0.40), ah - Inches(0.40))
        tf_ag = tb_ag.text_frame
        tf_ag.word_wrap = True

        p1 = tf_ag.paragraphs[0]
        p1.text = ag_t
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = C_BLUE_ACCENT
        p1.space_after = Pt(8)

        p2 = tf_ag.add_paragraph()
        p2.text = ag_d
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = C_TEXT_MUTED
        p2.line_spacing = 1.25

    # =========================================================================
    # SLIDE 3: PHẦN 1 - GIỚI THIỆU DỰ ÁN & LÝ DO CHỌN SPI FLASH NVM
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Phần 1: Giới Thiệu Dự Án", "Thiết Bị Kiểm Soát Ra Vào Offline & Lý Do Lựa Chọn SPI Flash", 3, total_slides=TOTAL_SLIDES)

    p1_cards = [
        ("TÍNH CẤP THIẾT CỦA HỆ THỐNG OFFLINE", C_BLUE_ACCENT, [
            ("Hạn chế của Access Control dựa trên Cloud:", "Phụ thuộc kết nối Internet, độ trễ mạng lớn (vài trăm ms đến vài giây), nguy cơ gián đoạn khi đứt cáp và lộ lọt dữ liệu định danh nhân sự."),
            ("Định hướng hệ thống Offline Standalone:", "SoC tự giải mã thẻ và tự đối chiếu phân quyền tại chỗ với độ trễ siêu nhỏ (< 10 µs), vận hành ổn định 24/7 tại các địa bàn biệt lập, bảo mật cao."),
            ("Nhu cầu tự chủ thiết kế vi mạch ASIC:", "Loại bỏ hoàn toàn nguy cơ backdoor phần cứng, làm chủ từ kiến trúc RTL đến quy trình chế tạo vi mạch bán dẫn.")
        ]),
        ("VAI TRÒ CỐT TỬ CỦA BỘ NHỚ SPI FLASH NVM", C_GREEN, [
            ("Lưu trữ bất biến không cần pin nuôi:", "Khác với RAM mất sạch dữ liệu khi ngắt điện, SPI Flash lưu trữ danh sách thẻ (Whitelist) và nhật ký quẹt thẻ (Access Logs) an toàn trên 20 năm."),
            ("Dung lượng lớn với chi phí tối thiểu:", "Dung lượng 32 Mbit (4MB) cho phép lưu tới 4,096 thẻ người dùng và 512 bản ghi nhật ký kiểm toán với chi phí linh kiện cực rẻ."),
            ("Tiết kiệm chân kết nối ngoại vi ASIC:", "Chuẩn SPI chỉ dùng 4 chân tín hiệu (CS, SCK, MOSI, MISO), tối ưu hóa số chân pad (chỉ 31 chân toàn vi mạch) và diện tích die silicon."),
            ("Hỗ trợ chế độ thực thi lệnh tại chỗ (XIP):", "Cho phép CPU PicoRV32 kéo trực tiếp mã lệnh từ Flash chạy mà không cần nạp vào SRAM lớn.")
        ])
    ]

    for i, (c_title, c_col, points) in enumerate(p1_cards):
        cx = Inches(0.8) + i * Inches(5.9)
        cw = Inches(5.7)
        add_card(s3, cx, Inches(1.35), cw, Inches(5.4), bg_color=C_WHITE, border_color=c_col, border_width=2.0)
        tb_c = s3.shapes.add_textbox(cx + Inches(0.2), Inches(1.5), cw - Inches(0.4), Inches(5.1))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        pt = tf_c.paragraphs[0]
        pt.text = c_title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = c_col
        pt.space_after = Pt(12)

        for p_lbl, p_val in points:
            p = tf_c.add_paragraph()
            p.text = f"• {p_lbl} {p_val}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(10)
            p.line_spacing = 1.2

    # =========================================================================
    # SLIDE 4: PHẦN 2 - LỰA CHỌN PHẦN CỨNG: BASYS 3 & RDM6300
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Phần 2: Nền Tảng Phần Cứng", "Lựa Chọn Nền Tảng Phần Cứng: FPGA Basys 3 & Đầu Đọc RFID RDM6300", 4, total_slides=TOTAL_SLIDES)

    hw_cards = [
        ("NỀN TẢNG FPGA DIGILENT BASYS 3", C_BLUE_ACCENT, [
            ("Chip FPGA Xilinx Artix-7 (XC7A35T):", "33,280 Logic Cells, 50 Block RAM 36Kb, 90 DSP Slices - nền tảng lý tưởng tạo mẫu SoC."),
            ("Tích hợp sẵn chip Flash SPI 32Mbit:", "Chip Spansion S25FL032P giao tiếp qua nguyên thủy STARTUPE2 của Xilinx."),
            ("Tích hợp mạch nạp & UART FTDI:", "Cổng micro-USB cung cấp nguồn và kênh UART nối tiếp kết nối trực tiếp với máy tính."),
            ("Hệ thống hiển thị & chẩn đoán tại chỗ:", "16 LED đơn, 16 công tắc gạt (Switch) và 4 nút nhấn phục vụ chẩn đoán trạng thái.")
        ]),
        ("MODULE ĐẦU ĐỌC RFID RDM6300 125KHZ", C_AMBER, [
            ("Tần số sóng mang 125 kHz EM4100:", "Tần số tiêu chuẩn công nghiệp cho thẻ cảm ứng tầm gần (khoảng cách đọc 20 - 50 mm)."),
            ("Giao tiếp UART 9600-8-N-1 tự động:", "Module tự động điều chế và phát ra khung dữ liệu ASCII 14 byte ngay khi phát hiện thẻ."),
            ("Cấu trúc khung dữ liệu an toàn:", "Gồm 1 byte Start (0x02), 10 byte mã thẻ Hex, 2 byte Checksum XOR và 1 byte Stop (0x03)."),
            ("Ăng-ten cuộn cảm rời linh hoạt:", "Dễ dàng lắp đặt tại cửa kiểm soát mà không gây can nhiễu từ trường đến mạch số.")
        ])
    ]

    for i, (hw_title, hw_col, points) in enumerate(hw_cards):
        hx = Inches(0.8) + i * Inches(5.9)
        hw = Inches(5.7)
        add_card(s4, hx, Inches(1.35), hw, Inches(5.4), bg_color=C_WHITE, border_color=hw_col, border_width=2.0)
        tb_h = s4.shapes.add_textbox(hx + Inches(0.2), Inches(1.5), hw - Inches(0.4), Inches(5.1))
        tf_h = tb_h.text_frame
        tf_h.word_wrap = True

        pt = tf_h.paragraphs[0]
        pt.text = hw_title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = hw_col
        pt.space_after = Pt(12)

        for p_lbl, p_val in points:
            p = tf_h.add_paragraph()
            p.text = f"• {p_lbl} {p_val}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(10)
            p.line_spacing = 1.2

    # =========================================================================
    # SLIDE 5: PHẦN 2 - LỰA CHỌN PHẦN MỀM: VIVADO, OPENLANE 2 & TOOLCHAIN
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Phần 2: Chuỗi Công Cụ Phát Triển", "Chuỗi Công Cụ Phát Triển: FPGA Vivado, ASIC OpenLane 2 & RISC-V GCC", 5, total_slides=TOTAL_SLIDES)

    sw_tools = [
        ("MÔ PHỎNG & FPGA: XILINX VIVADO", C_BLUE_ACCENT, [
            ("Môi trường Vivado 2025.1:", "Công cụ số 1 thế giới cho tổng hợp logic và nạp FPGA Artix-7."),
            ("Vivado Simulator (xsim):", "Mô phỏng chính xác mức chu kỳ clock cho cả testbench RTL thuần và Top SoC."),
            ("Tự động hóa bằng Batch Script:", "Tích hợp sẵn run_sim_uart.bat và run_sim_ping.bat chạy 1-click.")
        ]),
        ("THIẾT KẾ VẬT LÝ ASIC: OPENLANE 2", C_PURPLE, [
            ("Chuỗi công cụ nguồn mở OpenLane 2:", "Hiện thực hóa luồng RTL-to-GDSII tự động trên tiến trình SkyWater 130nm."),
            ("Công cụ chuẩn công nghiệp tích hợp:", "Yosys (Synthesis), OpenROAD (P&R), Magic & KLayout (DRC), Netgen (LVS)."),
            ("Sign-off Tape-out Ready:", "Tự động chèn diode triệt tiêu Antenna, tối ưu hóa timing MET TIMING 50 MHz.")
        ]),
        ("TOOLCHAIN FIRMWARE: RISC-V GCC & C", C_GREEN, [
            ("Bộ biên dịch riscv32-unknown-elf-gcc:", "Biên dịch mã C thành tập lệnh RV32I freestanding tối ưu kích thước."),
            ("Linker Script sections.lds:", "Định vị mã nguồn vào Flash XIP (0x0025_0000) và biến vào 1KB SRAM."),
            ("Host Console C tương tác:", "Phần mềm quản trị trên máy tính hỗ trợ 11 lệnh điều khiển qua UART.")
        ])
    ]

    for i, (tool_title, tool_col, points) in enumerate(sw_tools):
        tx = Inches(0.8) + i * Inches(3.95)
        add_card(s5, tx, Inches(1.4), Inches(3.75), Inches(5.3), bg_color=C_WHITE, border_color=tool_col, border_width=2.0)
        tb_t = s5.shapes.add_textbox(tx + Inches(0.18), Inches(1.6), Inches(3.39), Inches(4.9))
        tf_t = tb_t.text_frame
        tf_t.word_wrap = True

        pt = tf_t.paragraphs[0]
        pt.text = tool_title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = tool_col
        pt.space_after = Pt(12)

        for p_lbl, p_val in points:
            p = tf_t.add_paragraph()
            p.text = f"• {p_lbl} {p_val}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(8)
            p.line_spacing = 1.15

    # =========================================================================
    # SLIDE 6: PHẦN 3 - SƠ ĐỒ KHỐI KIẾN TRÚC SOC & 3 BƯỚC HIỆN THỰC CO-DESIGN
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Phần 3: Kiến Trúc Vi Hệ Thống SoC",
               "Sơ Đồ Khối Kiến Trúc Vi Hệ Thống SoC PicoRV32 & 3 Bước Hiện Thực Co-Design", 6, total_slides=TOTAL_SLIDES)

    # Bên trái: Ảnh Block Diagram Draw.io đen trắng chuẩn mực
    if os.path.exists(img_fig1):
        add_card(s6, Inches(0.8), Inches(1.35), Inches(7.4), Inches(5.4), bg_color=C_WHITE, border_color=C_BLUE_ACCENT, border_width=1.5)
        s6.shapes.add_picture(img_fig1, Inches(0.9), Inches(1.45), Inches(7.2), Inches(5.2))

    # Bên phải: 3 bước hiện thực Co-Design trực quan, cực kỳ dễ thuyết trình
    rx6 = Inches(8.4)
    rw6 = Inches(4.13)

    step_cards = [
        ("BƯỚC 1: CÁC KHỐI NỀN TẢNG CỐT LÕI (RTL CƠ SỞ)", C_BLUE_ACCENT, Inches(1.35), Inches(1.70), [
            ("Nhân CPU picorv32.v:", "RV32I 32-bit, giao tiếp bộ nhớ đồng bộ mem_valid & mem_ready."),
            ("Flash Controller spimemio.v:", "Điều khiển Flash SPI ngoài, hỗ trợ đọc lệnh trực tiếp XIP."),
            ("Khối data_sram (1KB):", "Bộ nhớ RAM nội bộ tốc độ cao, phản hồi 1 chu kỳ clock cho Stack & biến.")
        ]),
        ("BƯỚC 2: CẤU HÌNH HỆ THỐNG CHẠY FIRMWARE", C_AMBER, Inches(3.18), Inches(1.75), [
            ("Bus soc_interconnect.v:", "Giải mã địa chỉ 0-delay phân luồng: SRAM, Flash, RFID, UART, GPIO."),
            ("Khớp Boot Vector:", "PROGADDR_RESET (RTL) = FLASH ORIGIN (Linker) = 0x0025_0000."),
            ("Khởi tạo ngăn xếp:", "STACKADDR (RTL) = _stack_top (Linker) = 0x0400 nạp vào sp (start.s).")
        ]),
        ("BƯỚC 3: LẬP TRÌNH C ĐIỀU KHIỂN NGOẠI VI (MMIO)", C_GREEN, Inches(5.05), Inches(1.70), [
            ("REG_RFID_UART_DAT (0x10000004):", "Đọc không khóa mã thẻ 14-byte từ hàng đợi FIFO đệm 32B."),
            ("REG_PC_UART_DAT (0x30000004):", "Giao tiếp máy tính: nhận lệnh quản trị, xuất nhật ký kiểm toán."),
            ("REG_GPIO_LEDS (0x40000000):", "Điều khiển 16 LED trạng thái và kích hoạt relay đóng/mở chốt cửa.")
        ])
    ]

    for s_title, s_col, sy, sh, s_points in step_cards:
        add_card(s6, rx6, sy, rw6, sh, bg_color=C_WHITE, border_color=s_col, border_width=1.5)
        tb_s = s6.shapes.add_textbox(rx6 + Inches(0.12), sy + Inches(0.08), rw6 - Inches(0.24), sh - Inches(0.16))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = tf_s.margin_bottom = 0

        p_t = tf_s.paragraphs[0]
        p_t.text = s_title
        p_t.font.name = "Segoe UI"
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = s_col
        p_t.space_after = Pt(3)

        for p_lbl, p_val in s_points:
            p = tf_s.add_paragraph()
            p.text = f"• {p_lbl} {p_val}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(8.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(2)
            p.line_spacing = 1.10

    # =========================================================================
    # SLIDE 7: PHẦN 4 - HỆ THỐNG TESTBENCH: UART RTL (TB_UART_RTL.V)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Phần 4: Hệ Thống Testbench", "Testbench 1: Kiểm Thử Phần Cứng RTL Thuần Cho UART MMIO & FIFO (tb_uart_rtl.v)", 7, total_slides=TOTAL_SLIDES)

    lx7 = Inches(0.8)
    lw7 = Inches(5.7)
    add_card(s7, lx7, Inches(1.35), lw7, Inches(5.4), bg_color=C_WHITE, border_color=C_BLUE_ACCENT, border_width=1.5)
    tb_l7 = s7.shapes.add_textbox(lx7 + Inches(0.2), Inches(1.5), lw7 - Inches(0.4), Inches(5.1))
    tf_l7 = tb_l7.text_frame
    tf_l7.word_wrap = True

    pt7 = tf_l7.paragraphs[0]
    pt7.text = "5 KỊCH BẢN KIỂM THỬ UART RTL THUẦN"
    pt7.font.name = "Segoe UI"
    pt7.font.size = Pt(13)
    pt7.font.bold = True
    pt7.font.color.rgb = C_BLUE_ACCENT
    pt7.space_after = Pt(10)

    tb_uart_scenarios = [
        ("Mục tiêu:", "Xác minh tính đúng đắn phần cứng của uart_mmio.v, sync_fifo.v và simpleuart.v mà không cần CPU."),
        ("1. Default Divider:", "Kiểm tra thanh ghi Prescaler tại offset 0x00 nạp đúng giá trị mặc định 5208 (9600 baud @ 50MHz)."),
        ("2. Divider Reconfig:", "Ghi giá trị chia tần mới và đọc lại qua bus MMIO, xác nhận mạch thanh ghi hoạt động chuẩn xác."),
        ("3. Serial TX Waveform:", "Ghi ký tự vào offset 0x04, kiểm tra dạng sóng nối tiếp tx_o (Start bit = 0, 8 data bits LSB-first, Stop bit = 1)."),
        ("4. Serial RX & FIFO:", "Bắn chuỗi bit vào rx_i, cờ rx_activity_o tích cực, kiểm tra dữ liệu nạp vào FIFO và trả về 0xFFFFFFFF khi rỗng."),
        ("5. Multi-Byte Burst:", "Bắn liên tiếp chuỗi byte kiểm tra cơ chế chống tràn của hàng đợi FIFO 32 byte.")
    ]

    for p_lbl, p_val in tb_uart_scenarios:
        p = tf_l7.add_paragraph()
        p.text = f"• {p_lbl} {p_val}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(8)
        p.line_spacing = 1.15

    rx7 = Inches(6.8)
    rw7 = Inches(5.73)
    code_lines_7 = [
        "Vivado Simulator v2025.1 - xelab & xsim",
        "Analyzing Verilog file rtl/uart/sync_2ff.v",
        "Analyzing Verilog file rtl/uart/sync_fifo.v",
        "Analyzing Verilog file rtl/uart/simpleuart.v",
        "Analyzing Verilog file rtl/uart/uart_mmio.v",
        "Analyzing Verilog file tb/tb_uart_rtl.v",
        "",
        "[TEST 1] Checking Default Prescaler Divider...",
        "  [PASS] Prescaler default matches DEFAULT_DIV = 5208",
        "[TEST 2] Reconfiguring Divider to 10416 (4800 baud)...",
        "  [PASS] Read-back matches written value 10416",
        "[TEST 3] Transmitting Byte 0x55 ('U') over TX...",
        "  [PASS] TX Waveform matches 9600-8-N-1 timing!",
        "[TEST 4] Receiving Serial Byte 0xA5 into FIFO...",
        "  [PASS] rx_activity_o asserted! Read 0xA5 cleanly!",
        "  [PASS] FIFO Empty flag returns 0xFFFFFFFF",
        "[TEST 5] Multi-Byte Burst Test (16 bytes)...",
        "  [PASS] All 16 bytes drained without data corruption!"
    ]
    add_code_box(s7, rx7, Inches(1.35), rw7, Inches(5.4), "tb/run_sim_uart.bat output", code_lines_7, "=== UART RTL TESTBENCH: 100% PASS ===")

    # =========================================================================
    # SLIDE 8: PHẦN 4 - HỆ THỐNG TESTBENCH: TOP SOC BOOT & PING (TB_UART_PING.V)
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Phần 4: Hệ Thống Testbench", "Testbench 2: Kiểm Thử Tích Hợp Toàn Diện Top SoC Boot Flash & Ping (tb_uart_ping.v)", 8, total_slides=TOTAL_SLIDES)

    lx8 = Inches(0.8)
    lw8 = Inches(5.7)
    add_card(s8, lx8, Inches(1.35), lw8, Inches(5.4), bg_color=C_WHITE, border_color=C_GREEN, border_width=1.5)
    tb_l8 = s8.shapes.add_textbox(lx8 + Inches(0.2), Inches(1.5), lw8 - Inches(0.4), Inches(5.1))
    tf_l8 = tb_l8.text_frame
    tf_l8.word_wrap = True

    pt8 = tf_l8.paragraphs[0]
    pt8.text = "5 KỊCH BẢN KIỂM THỬ TÍCH HỢP TOP SOC"
    pt8.font.name = "Segoe UI"
    pt8.font.size = Pt(13)
    pt8.font.bold = True
    pt8.font.color.rgb = C_GREEN
    pt8.space_after = Pt(10)

    tb_soc_scenarios = [
        ("Mục tiêu:", "Xác minh hệ thống Top-level hoàn chỉnh gồm CPU PicoRV32, spimemio, 1KB SRAM, Interconnect và UART."),
        ("1. Power-On Reset & Boot Flash:", "PicoRV32 thức dậy tại 0x0025_0000, spimemio kéo từng từ lệnh mã C từ firmware.hex qua XIP."),
        ("2. C Startup Banner:", "CPU thực thi mã C trong main.c, khởi tạo ngoại vi và in toàn bộ chuỗi chào mừng ra UART."),
        ("3. Host Ping Processing:", "Testbench đóng vai trò PC Host gửi byte lệnh 'P' (0x50) và '\\n' (0x0A) qua cổng nối tiếp."),
        ("4. Response Verification:", "CPU phản hồi chuỗi 'PONG: PicoRV32 Active', testbench so khớp từng ký tự."),
        ("5. CPU Health & Zero-Trap:", "Khẳng định tín hiệu cpu_trap == 0 xuyên suốt quá trình chạy, không bị illegal instruction hay tràn RAM.")
    ]

    for p_lbl, p_val in tb_soc_scenarios:
        p = tf_l8.add_paragraph()
        p.text = f"• {p_lbl} {p_val}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(8)
        p.line_spacing = 1.15

    rx8 = Inches(6.8)
    rw8 = Inches(5.73)
    code_lines_8 = [
        "Vivado Simulator v2025.1 - xelab & xsim",
        "Analyzing Verilog file rtl/rdm6300_picorv32_soc.v",
        "Analyzing Verilog file tb/tb_uart_ping.v",
        "",
        "[BOOT] Releasing CPU Reset...",
        "[BOOT] PicoRV32 woke up at 0x00250000 (Flash XIP)",
        "[EXEC] Loading instructions from firmware.hex...",
        "[UART] Detected Startup Banner:",
        "  ================================================",
        "    RDM6300 PICORV32 SOC ACCESS CONTROLLER READY  ",
        "  ================================================",
        "[HOST] Sending Host Command: 'P' (Ping)...",
        "[RECV] CPU Response received via UART:",
        "  PONG: PicoRV32 Active",
        "[VERIFY] String matches expected response!",
        "[ASSERT] Checking cpu_trap signal...",
        "  cpu_trap == 0 (CONFIRMED: ZERO CPU TRAPS!)"
    ]
    add_code_box(s8, rx8, Inches(1.35), rw8, Inches(5.4), "tb/run_sim_ping.bat output", code_lines_8, "=== TOP SOC PING TESTBENCH: 100% PASS ===")

    # =========================================================================
    # SLIDE 9: PHẦN 5 - DEMO THỰC NGHIỆM TRÊN FPGA BASYS 3: THIẾT LẬP KẾT NỐI
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Phần 5: Demo Chức Năng Sản Phẩm", "Kiểm Chứng Thực Nghiệm Trên FPGA Basys 3: Thiết Lập Kết Nối Phần Cứng", 9, total_slides=TOTAL_SLIDES)

    demo_setup_cards = [
        ("KẾT NỐI ĐẦU ĐỌC RFID RDM6300", C_AMBER, [
            ("Cổng PMOD JA trên Basys 3:", "Cắm module RDM6300 trực tiếp vào cổng PMOD JA."),
            ("Chân JA1 (Pin J1 FPGA):", "Nhận luồng dữ liệu UART TX 9600-8-N-1 từ RDM6300."),
            ("Nguồn cấp 5V & GND:", "Lấy trực tiếp từ chân VCC/GND của cổng PMOD bo mạch."),
            ("Ăng-ten cuộn cảm rời:", "Đặt cách ly chống nhiễu từ trường với các linh kiện số trên bo mạch.")
        ]),
        ("KẾT NỐI MÁY TÍNH HOST PC", C_BLUE_ACCENT, [
            ("Cáp Micro-USB tiêu chuẩn:", "Vừa cấp nguồn +5V DC vừa mở cổng COM ảo qua chip FTDI."),
            ("Tốc độ baud 9600 bps:", "Giao tiếp Full-Duplex tin cậy giữa máy tính và SoC PicoRV32."),
            ("Phần mềm Host Console C:", "Giao diện dòng lệnh tương tác 11 chức năng quản trị cơ sở dữ liệu thẻ."),
            ("Trích xuất dữ liệu tự động:", "Tự động sao lưu lịch sử quẹt thẻ ra file CSV trên máy tính.")
        ]),
        ("ĐIỀU KHIỂN FLASH & HIỂN THỊ LED", C_PURPLE, [
            ("Nguyên thủy STARTUPE2:", "Đưa xung nhịp Flash SCK vào USRCCLKO để làm chủ chân CCLK nối tới chip Flash."),
            ("Chip Flash Spansion S25FL032P:", "Lưu trữ bền vững danh mục Whitelist (Sector 48) và Access Logs (Sector 49)."),
            ("LED[0] (Heartbeat):", "Nhấp nháy 1 Hz báo hiệu CPU PicoRV32 đang hoạt động bình thường."),
            ("LED[1] & LED[2]:", "LED[1] sáng khi chip Flash đang bận ghi; LED[2] sáng khi quẹt thẻ hợp lệ.")
        ])
    ]

    for i, (ds_title, ds_col, points) in enumerate(demo_setup_cards):
        dx = Inches(0.8) + i * Inches(3.95)
        add_card(s9, dx, Inches(1.4), Inches(3.75), Inches(5.3), bg_color=C_WHITE, border_color=ds_col, border_width=2.0)
        tb_d = s9.shapes.add_textbox(dx + Inches(0.18), Inches(1.6), Inches(3.39), Inches(4.9))
        tf_d = tb_d.text_frame
        tf_d.word_wrap = True

        pt = tf_d.paragraphs[0]
        pt.text = ds_title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = ds_col
        pt.space_after = Pt(12)

        for p_lbl, p_val in points:
            p = tf_d.add_paragraph()
            p.text = f"• {p_lbl} {p_val}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(8)
            p.line_spacing = 1.15

    # =========================================================================
    # SLIDE 10: PHẦN 5 - DEMO THỰC NGHIỆM TRÊN FPGA BASYS 3: CÁC KỊCH BẢN THỰC TẾ
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Phần 5: Demo Chức Năng Sản Phẩm", "Các Kịch Bản Thực Nghiệm Quẹt Thẻ Thực Tế Trên FPGA Basys 3", 10, total_slides=TOTAL_SLIDES)

    demo_cases = [
        ("KỊCH BẢN 1: PING PHẦN CỨNG", C_BLUE_ACCENT, [
            ("Thao tác:", "Chọn menu [1] trên Host Console gửi lệnh 'P'."),
            ("Kết quả:", "PicoRV32 phản hồi ngay 'PONG: PicoRV32 Active'."),
            ("Ý nghĩa:", "Khẳng định CPU, Interconnect và UART hoạt động thông suốt.")
        ]),
        ("KỊCH BẢN 2: THẺ CHƯA ĐĂNG KÝ", C_ROSE, [
            ("Thao tác:", "Quẹt thẻ thật 0007508976 chưa có trong Flash."),
            ("Phản hồi:", "Hệ thống bóc tách UID 00007293F0 -> 'ACCESS:DENIED'."),
            ("Tác động:", "Nháy LED đỏ cảnh báo, ghi vết FAIL vào Flash Sector 49.")
        ]),
        ("KỊCH BẢN 3: ĐĂNG KÝ THẺ VÀO FLASH", C_AMBER, [
            ("Thao tác:", "Chọn menu [2] nhập 10 số in trên thẻ (0007508976)."),
            ("Xử lý:", "CPU nạp flashio_worker vào SRAM ghi vào Sector 48."),
            ("Kết quả:", "Lưu thành công tại Slot #0 (0x300000), nháy LED Flash.")
        ]),
        ("KỊCH BẢN 4: QUẸT THẺ HỢP LỆ", C_GREEN, [
            ("Thao tác:", "Quẹt lại thẻ 0007508976 lên đầu đọc RDM6300."),
            ("Phản hồi:", "SoC tra cứu trúng Slot #0 -> 'ACCESS:GRANTED'."),
            ("Tác động:", "Bật sáng LED xanh mở cửa, ghi vết SUCC vào Flash Sector 49.")
        ])
    ]

    for i, (dc_title, dc_col, points) in enumerate(demo_cases):
        col = i % 2
        row = i // 2
        cx = Inches(0.8) + col * Inches(5.9)
        cy = Inches(1.4) + row * Inches(2.65)
        cw = Inches(5.6)
        ch = Inches(2.35)

        add_card(s10, cx, cy, cw, ch, bg_color=C_WHITE, border_color=dc_col, border_width=2.0)
        tb_c = s10.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.15), cw - Inches(0.4), ch - Inches(0.3))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True

        pt = tf_c.paragraphs[0]
        pt.text = dc_title
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12.5)
        pt.font.bold = True
        pt.font.color.rgb = dc_col
        pt.space_after = Pt(8)

        for p_lbl, p_val in points:
            p = tf_c.add_paragraph()
            p.text = f"• {p_lbl} {p_val}"
            p.font.name = "Segoe UI"
            p.font.size = Pt(10.5)
            p.font.color.rgb = C_TEXT_DARK
            p.space_after = Pt(4)
            p.line_spacing = 1.15

    # =========================================================================
    # SLIDE 11: PHẦN 6 - THIẾT KẾ VẬT LÝ ASIC: LUỒNG OPENLANE 2 (SKYWATER 130NM)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Phần 6: Thiết Kế Vật Lý ASIC", "Quy Trình Thiết Kế Vật Lý Vi Mạch RTL-to-GDSII Trên OpenLane 2", 11, total_slides=TOTAL_SLIDES)

    lx11 = Inches(0.8)
    lw11 = Inches(5.7)
    code_lines_11 = [
        '{',
        '  "DESIGN_NAME": "rdm6300_picorv32_soc",',
        '  "PDK": "sky130A",',
        '  "STD_CELL_LIBRARY": "sky130_fd_sc_hd",',
        '  "CLOCK_PORT": "clk",',
        '  "CLOCK_PERIOD": 20.0, // 50 MHz',
        '  "FP_CORE_UTIL": 26,   // 26% density',
        '  "MAX_FANOUT_CONSTRAINT": 12,',
        '  "PL_RESIZER_HOLD_SLACK_MARGIN": 0.6,',
        '  "RUN_HEURISTIC_DIODE_INSERTION": true,',
        '  "HEURISTIC_ANTENNA_THRESHOLD": 24,',
        '  "GRT_ANTENNA_ITERS": 35,',
        '  "RUN_ANTENNA_REPAIR": true,',
        '  "SYNTH_STRATEGY": "AREA 0"',
        '}'
    ]
    add_code_box(s11, lx11, Inches(1.35), lw11, Inches(5.4), "Cấu hình config.json chính thức", code_lines_11)

    rx11 = Inches(6.8)
    rw11 = Inches(5.73)
    add_card(s11, rx11, Inches(1.35), rw11, Inches(5.4), bg_color=C_WHITE, border_color=C_ROSE, border_width=2.0)
    tb_r11 = s11.shapes.add_textbox(rx11 + Inches(0.2), Inches(1.5), rw11 - Inches(0.4), Inches(5.1))
    tf_r11 = tb_r11.text_frame
    tf_r11.word_wrap = True

    pr11 = tf_r11.paragraphs[0]
    pr11.text = "CÁC CHIẾN LƯỢC VẬT LÝ THEN CHỐT"
    pr11.font.name = "Segoe UI"
    pr11.font.size = Pt(13)
    pr11.font.bold = True
    pr11.font.color.rgb = C_ROSE
    pr11.space_after = Pt(10)

    asic_strategies = [
        ("Mật độ diện tích lõi 26% (FP_CORE_UTIL):", "Khởi tạo mật độ cell 26%, để lại 74% diện tích cho các kênh định tuyến kim loại và chèn diode bảo vệ plasma."),
        ("Ràng buộc xung nhịp 50 MHz (CLOCK_PERIOD = 20.0 ns):", "Ràng buộc định thời thống nhất từ khâu tổng hợp logic Yosys, xây dựng cây clock CTS tới khâu phân tích định thời tĩnh STA."),
        ("Chèn Diode phỏng đoán (HEURISTIC_ANTENNA_THRESHOLD = 24):", "Thuật toán tối ưu tự động chèn diode tiêu tán điện tích plasma khi tỷ lệ dây/cực cổng vượt ngưỡng 24, triệt tiêu 100% lỗi Antenna."),
        ("Biên an toàn Hold Guard-band 0.60 ns:", "Đảm bảo bộ tối ưu hóa tế bào Resizer xử lý sạch sẽ các vi phạm chạy đua dữ liệu (Hold Violations) sau bước đi dây chi tiết.")
    ]

    for p_lbl, p_val in asic_strategies:
        p = tf_r11.add_paragraph()
        p.text = f"• {p_lbl} {p_val}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = C_TEXT_DARK
        p.space_after = Pt(8)
        p.line_spacing = 1.15

    # =========================================================================
    # SLIDE 12: PHẦN 6 - THIẾT KẾ VẬT LÝ ASIC: BẢN VẼ LAYOUT & KÝ DUYỆT SIGN-OFF
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Phần 6: Thiết Kế Vật Lý ASIC", "Bản Vẽ Layout OpenROAD & Báo Cáo Ký Duyệt Sign-off Toàn Diện", 12, total_slides=TOTAL_SLIDES)

    if os.path.exists(img_openroad):
        add_card(s12, Inches(0.8), Inches(1.35), Inches(5.7), Inches(4.3), bg_color=C_WHITE, border_color=C_BLUE_ACCENT, border_width=1.5)
        s12.shapes.add_picture(img_openroad, Inches(0.9), Inches(1.45), Inches(5.5), Inches(4.1))

    if os.path.exists(img_signoff):
        add_card(s12, Inches(6.8), Inches(1.35), Inches(5.73), Inches(4.3), bg_color=C_WHITE, border_color=C_GREEN, border_width=1.5)
        s12.shapes.add_picture(img_signoff, Inches(6.9), Inches(1.45), Inches(5.53), Inches(4.1))

    add_card(s12, Inches(0.8), Inches(5.8), Inches(11.733), Inches(1.05), bg_color=C_WHITE, border_color=C_CARD_BORDER, border_width=1.5)
    tb_c12 = s12.shapes.add_textbox(Inches(0.95), Inches(5.85), Inches(11.4), Inches(0.95))
    tf_c12 = tb_c12.text_frame
    tf_c12.word_wrap = True

    p_c12_1 = tf_c12.paragraphs[0]
    p_c12_1.text = "KẾT QUẢ KÝ DUYỆT TAPE-OUT READY: 0 ANTENNA | 0 LVS | 0 DRC | MET TIMING 50 MHZ"
    p_c12_1.font.name = "Segoe UI"
    p_c12_1.font.size = Pt(11)
    p_c12_1.font.bold = True
    p_c12_1.font.color.rgb = C_GREEN
    p_c12_1.space_after = Pt(2)

    p_c12_2 = tf_c12.add_paragraph()
    p_c12_2.text = "• Bố cục vật lý: Lưới nguồn PDN met4/met5 dày đặc, phân bổ 31 chân pad đối xứng trên 4 cạnh die, độ sụt áp IR drop < 1.5% an toàn.\n• Định thời & Ký duyệt: WNS >= 0.00 ns tại 50 MHz trên cả 9 góc đo công nghệ từ -40°C đến 100°C; Netgen xác nhận 100% khớp netlist RTL."
    p_c12_2.font.name = "Segoe UI"
    p_c12_2.font.size = Pt(9.5)
    p_c12_2.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 13: TỔNG KẾT ĐỀ TÀI & HƯỚNG PHÁT TRIỂN (THANK YOU SLIDE)
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    bg13 = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg13.fill.solid()
    bg13.fill.fore_color.rgb = C_NAVY_DARK
    bg13.line.fill.background()

    card13 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    card13.fill.solid()
    card13.fill.fore_color.rgb = C_NAVY_MID
    card13.line.color.rgb = C_BLUE_ACCENT
    card13.line.width = Pt(2.0)

    tb_t13 = s13.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(10.9), Inches(5.3))
    tf_t13 = tb_t13.text_frame
    tf_t13.word_wrap = True

    p_end1 = tf_t13.paragraphs[0]
    p_end1.text = "TỔNG KẾT ĐỒ ÁN VÀ CÁC ĐÓNG GÓP NỔI BẬT"
    p_end1.font.name = "Segoe UI"
    p_end1.font.size = Pt(22)
    p_end1.font.bold = True
    p_end1.font.color.rgb = C_CYAN_ACCENT
    p_end1.space_after = Pt(16)

    contributions = [
        ("1. Tự chủ thiết kế kiến trúc SoC:", "Xây dựng hoàn chỉnh vi mạch tích hợp CPU RISC-V PicoRV32, cầu nối bus trung tâm soc_interconnect.v, 1KB SRAM, bộ điều khiển SPI Flash spimemio.v và ngoại vi UART FIFO."),
        ("2. Tối ưu hóa kiến trúc thực thi nhúng:", "Hiện thực hóa cơ chế thực thi XIP từ SPI Flash kết hợp nạp động routine flashio_worker vào 1KB SRAM để ghi/xóa cơ sở dữ liệu Whitelist và nhật ký Access Log không xung đột bus."),
        ("3. Kiểm chứng thực nghiệm 100%:", "Xây dựng 2 testbench mô phỏng tinh gọn (tb_uart_rtl.v, tb_uart_ping.v) và kiểm chứng thành công trên bo mạch FPGA Basys 3 với thẻ RFID thật và đầu đọc RDM6300 thật."),
        ("4. Đạt chuẩn xuất xưởng ASIC (Tape-out Ready):", "Hoàn thành quy trình thiết kế vật lý trên OpenLane 2 với tiến trình SkyWater 130nm: 0 vi phạm Antenna, 0 lỗi LVS, 0 lỗi DRC, MET TIMING ở tần số 50 MHz trên cả 9 góc đo công nghệ.")
    ]

    for c_title, c_desc in contributions:
        p = tf_t13.add_paragraph()
        p.text = f"{c_title} {c_desc}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.space_after = Pt(8)
        p.line_spacing = 1.15

    p_ty = tf_t13.add_paragraph()
    p_ty.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY CÔ VÀ HỘI ĐỒNG ĐÃ LẮNG NGHE!"
    p_ty.font.name = "Segoe UI"
    p_ty.font.size = Pt(15)
    p_ty.font.bold = True
    p_ty.font.color.rgb = C_WHITE
    p_ty.space_before = Pt(16)

    # -------------------------------------------------------------
    # LƯU FILE VÀ COPY
    # -------------------------------------------------------------
    out_path = os.path.join(cur_dir, "Bao_Cao_Do_An_RDM6300_PicoRV32_SoC.pptx")
    out_v2   = os.path.join(cur_dir, "Bao_Cao_Do_An_RDM6300_PicoRV32_SoC_v2.pptx")
    prs.save(out_path)
    shutil.copy2(out_path, out_v2)
    print(f"Presentation generated successfully: {out_path} ({os.path.getsize(out_path)} bytes)")

    doc_out    = os.path.join(doc_dir, "Bao_Cao_Do_An_RDM6300_PicoRV32_SoC.pptx")
    doc_out_v2 = os.path.join(doc_dir, "Bao_Cao_Do_An_RDM6300_PicoRV32_SoC_v2.pptx")
    shutil.copy2(out_path, doc_out)
    shutil.copy2(out_path, doc_out_v2)
    print(f"Copied presentation to document root: {doc_out} & {doc_out_v2}")

if __name__ == "__main__":
    create_deck()
