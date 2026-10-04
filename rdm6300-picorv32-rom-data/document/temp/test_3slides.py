# -*- coding: utf-8 -*-
"""
test_3slides.py: Tạo 3 slide cấu hình code theo đúng yêu cầu:
- Show các dòng code ở trên
- Đoạn giải thích ở dưới
- Nối các file với nhau
- Dùng phong cách Draw.io ĐEN TRẮNG hoàn toàn
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_test():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Bảng màu ĐEN TRẮNG chuẩn Draw.io
    C_BLACK      = RGBColor(0, 0, 0)         # #000000
    C_WHITE      = RGBColor(255, 255, 255)   # #FFFFFF
    C_BG_PAGE    = RGBColor(255, 255, 255)   # #FFFFFF Nền trắng chuẩn Draw.io
    C_GRAY_LIGHT = RGBColor(248, 250, 252)   # #F8FAFC
    C_GRAY_LINE  = RGBColor(203, 213, 225)   # #CBD5E1
    C_TEXT_DARK  = RGBColor(0, 0, 0)         # #000000
    C_TEXT_MUTED = RGBColor(71, 85, 105)     # #475569

    TOTAL_SLIDES = 17

    def add_header(slide, category, title, slide_num):
        # Nền slide trắng chuẩn Draw.io
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG_PAGE
        bg.line.fill.background()

        # Thanh viền đen trên cùng kiểu tài liệu kỹ thuật
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.18), Inches(11.733), Inches(0.04))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = C_BLACK
        top_bar.line.fill.background()

        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.24), Inches(11.733), Inches(0.90))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_BLACK
        p_cat.space_after = Pt(2)

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(17.5)
        p_title.font.bold = True
        p_title.font.color.rgb = C_BLACK

        # Footer
        footer_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.02))
        footer_line.fill.solid()
        footer_line.fill.fore_color.rgb = C_BLACK
        footer_line.line.fill.background()

        tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.10), Inches(8.5), Inches(0.3))
        tf_foot = tb_foot.text_frame
        tf_foot.margin_left = tf_foot.margin_top = tf_foot.margin_right = tf_foot.margin_bottom = 0
        p_foot = tf_foot.paragraphs[0]
        p_foot.text = "Đồ Án Tốt Nghiệp: SoC PicoRV32 RFID RDM6300 & SPI Flash | SV: Thái Tuấn Hiệp | GVHD: ThS. Nguyễn Văn Đông"
        p_foot.font.name = "Segoe UI"
        p_foot.font.size = Pt(9.5)
        p_foot.font.color.rgb = C_TEXT_MUTED

        tb_num = slide.shapes.add_textbox(Inches(11.0), Inches(7.10), Inches(1.533), Inches(0.3))
        tf_num = tb_num.text_frame
        tf_num.margin_left = tf_num.margin_right = tf_num.margin_top = tf_num.margin_bottom = 0
        p_num = tf_num.paragraphs[0]
        p_num.text = f"{slide_num:02d} / {TOTAL_SLIDES:02d}"
        p_num.alignment = PP_ALIGN.RIGHT
        p_num.font.name = "Segoe UI"
        p_num.font.size = Pt(9.5)
        p_num.font.bold = True
        p_num.font.color.rgb = C_BLACK

    # Helper: Khung file chuẩn Draw.io Đen Trắng (Code ở trên, Giải thích ở dưới)
    def add_drawio_file_box(slide, left, top, width, height, file_name, code_lines, explain_bullets, code_ratio=0.48):
        # 1. Khung ngoài hình chữ nhật Draw.io đen trắng
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = C_WHITE
        card.line.color.rgb = C_BLACK
        card.line.width = Pt(1.5)

        # 2. Thanh tiêu đề Header Draw.io (Đen chữ trắng)
        header_h = Inches(0.32)
        h_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, header_h)
        h_box.fill.solid()
        h_box.fill.fore_color.rgb = C_BLACK
        h_box.line.color.rgb = C_BLACK
        h_box.line.width = Pt(1.5)
        tf_h = h_box.text_frame
        tf_h.margin_left = Inches(0.12)
        tf_h.margin_top = Inches(0.04)
        p_h = tf_h.paragraphs[0]
        p_h.text = f"FILE: {file_name}"
        p_h.font.name = "Consolas"
        p_h.font.size = Pt(8.5)
        p_h.font.bold = True
        p_h.font.color.rgb = C_WHITE

        # 3. Phần code ở trên (CÁC DÒNG CODE Ở TRÊN)
        avail_h = height - header_h
        code_h = avail_h * code_ratio
        code_top = top + header_h

        # Nền code xám cực nhạt phân cách vùng code
        c_bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.04), code_top + Inches(0.04), width - Inches(0.08), code_h - Inches(0.06))
        c_bg.fill.solid()
        c_bg.fill.fore_color.rgb = RGBColor(248, 250, 252) # #F8FAFC
        c_bg.line.color.rgb = RGBColor(226, 232, 240)
        c_bg.line.width = Pt(0.8)

        tb_c = slide.shapes.add_textbox(left + Inches(0.08), code_top + Inches(0.05), width - Inches(0.16), code_h - Inches(0.08))
        tf_c = tb_c.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0

        for li, line in enumerate(code_lines):
            p = tf_c.paragraphs[0] if li == 0 else tf_c.add_paragraph()
            p.text = line
            p.font.name = "Consolas"
            p.font.size = Pt(7.4)
            p.font.bold = True
            p.font.color.rgb = C_BLACK
            p.line_spacing = 1.05
            p.space_after = 0

        # 4. Đường phân cách ngang giữa Code và Giải thích (Draw.io horizontal divider)
        div_y = code_top + code_h
        div_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, div_y, width, Inches(0.015))
        div_line.fill.solid()
        div_line.fill.fore_color.rgb = C_BLACK
        div_line.line.fill.background()

        # 5. Phần giải thích ở dưới (ĐOẠN GIẢI THÍCH Ở DƯỚI)
        exp_top = div_y + Inches(0.05)
        exp_h = avail_h - code_h - Inches(0.08)

        tb_e = slide.shapes.add_textbox(left + Inches(0.10), exp_top, width - Inches(0.20), exp_h)
        tf_e = tb_e.text_frame
        tf_e.word_wrap = True
        tf_e.margin_left = tf_e.margin_right = tf_e.margin_top = tf_e.margin_bottom = 0

        p_lbl = tf_e.paragraphs[0]
        p_lbl.text = "GIẢI THÍCH THIẾT KẾ & CO-DESIGN:"
        p_lbl.font.name = "Segoe UI"
        p_lbl.font.size = Pt(7.6)
        p_lbl.font.bold = True
        p_lbl.font.color.rgb = C_BLACK
        p_lbl.space_after = Pt(2)
        p_lbl.space_before = 0

        for bi, (b_title, b_desc) in enumerate(explain_bullets):
            p_b = tf_e.add_paragraph()
            p_b.space_before = 0
            p_b.space_after = Pt(1.5)
            p_b.line_spacing = 1.06

            r_t = p_b.add_run()
            r_t.text = f"• {b_title}: "
            r_t.font.name = "Segoe UI"
            r_t.font.size = Pt(7.4)
            r_t.font.bold = True
            r_t.font.color.rgb = C_BLACK

            r_d = p_b.add_run()
            r_d.text = b_desc
            r_d.font.name = "Segoe UI"
            r_d.font.size = Pt(7.4)
            r_d.font.color.rgb = RGBColor(30, 41, 59)

    # Helper: Mũi tên đen trắng kết nối giữa các file (Draw.io B&W Connector)
    def add_drawio_connector(slide, x, y, w, h, direction="right", label=None):
        shape_type = MSO_SHAPE.RIGHT_ARROW if direction == "right" else \
                     MSO_SHAPE.LEFT_ARROW if direction == "left" else \
                     MSO_SHAPE.DOWN_ARROW if direction == "down" else MSO_SHAPE.UP_ARROW
        arrow = slide.shapes.add_shape(shape_type, x, y, w, h)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = C_BLACK
        arrow.line.fill.background()

        if label:
            tb_lbl = slide.shapes.add_textbox(x - Inches(0.05), y - Inches(0.22), w + Inches(0.10), Inches(0.20))
            tf_lbl = tb_lbl.text_frame
            tf_lbl.word_wrap = False
            tf_lbl.margin_left = tf_lbl.margin_right = tf_lbl.margin_top = tf_lbl.margin_bottom = 0
            p_l = tf_lbl.paragraphs[0]
            p_l.alignment = PP_ALIGN.CENTER
            p_l.text = label
            p_l.font.name = "Segoe UI"
            p_l.font.size = Pt(6.5)
            p_l.font.bold = True
            p_l.font.color.rgb = C_BLACK

    # Helper: Khung tổng kết Co-Design ở chân slide (Đen Trắng Draw.io)
    def add_bottom_summary_box(slide, title, bullets):
        bot_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.32), Inches(11.733), Inches(0.68))
        bot_box.fill.solid()
        bot_box.fill.fore_color.rgb = C_WHITE
        bot_box.line.color.rgb = C_BLACK
        bot_box.line.width = Pt(1.5)

        tf_b = bot_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.12)
        tf_b.margin_top = Inches(0.04)

        p1 = tf_b.paragraphs[0]
        p1.text = f"■  CƠ CHẾ CO-DESIGN ĐỒNG BỘ: {title}"
        p1.font.name = "Segoe UI"
        p1.font.size = Pt(8.3)
        p1.font.bold = True
        p1.font.color.rgb = C_BLACK

        p2 = tf_b.add_paragraph()
        p2.text = bullets
        p2.font.name = "Segoe UI"
        p2.font.size = Pt(7.6)
        p2.font.color.rgb = C_BLACK
        p2.line_spacing = 1.08

    # =========================================================================
    # SLIDE 1 (Slide 07): KHỞI ĐỘNG CPU, FLASH XIP & NGĂN XẾP SRAM
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Phần 5: Cấu Hình Phần Cứng & Khởi Động Firmware (Slide 1/3)",
               "Co-Design Khởi Động: Cấu Hình CPU PicoRV32, Flash XIP & Ngăn Xếp SRAM", 7)

    # Box 1: rdm6300_picorv32_soc.v (Top Left)
    code_soc = [
        "// rtl/rdm6300_picorv32_soc.v - Tham số Top SoC PicoRV32",
        "parameter [31:0] PROGADDR_RESET = 32'h0025_0000; // Reset Vector XIP",
        "parameter [31:0] STACKADDR      = 32'h0000_0400; // Đỉnh stack 1KB SRAM",
        "parameter [0:0]  COMPRESSED_ISA = 1'b0;          // RV32I Base ISA",
        "wire mem_valid, mem_ready; wire [31:0] mem_addr, mem_wdata, mem_rdata;"
    ]
    exp_soc = [
        ("PROGADDR_RESET = 0x0025_0000", "CPU nhả reset nạp opcode trực tiếp từ Flash ngoài qua XIP."),
        ("STACKADDR = 0x0000_0400", "Đặt đỉnh ngăn xếp tại 1KB SRAM nội bộ, phản hồi 1 chu kỳ clock."),
        ("Native Mem Bus", "mem_valid & mem_ready bắt tay đồng bộ không qua cầu chuyển đổi.")
    ]
    add_drawio_file_box(s7, Inches(0.8), Inches(1.22), Inches(5.50), Inches(2.45),
                        "rtl/rdm6300_picorv32_soc.v [RTL Top Module]", code_soc, exp_soc, code_ratio=0.52)

    # Box 2: picorv32.v (Bottom Left)
    code_cpu = [
        "// rtl/core/picorv32.v - Nhân CPU RISC-V PicoRV32",
        "module picorv32 #(parameter [31:0] PROGADDR_RESET, STACKADDR, ...);",
        "output reg        mem_valid; input mem_ready;",
        "output reg [31:0] mem_addr, mem_wdata; input [31:0] mem_rdata;",
        "output reg        cpu_trap; // Tích cực cao khi lỗi instruction/bus"
    ]
    exp_cpu = [
        ("Cấu hình lõi Master", "Nhận tham số PROGADDR_RESET và STACKADDR để nạp vào PC và SP."),
        ("Native Memory Interface", "Phát mem_valid và mem_addr kéo opcode lệnh ngay chu kỳ đầu."),
        ("Bảo vệ hệ thống", "Ngõ ra cpu_trap cảnh báo khi xảy ra truy xuất bất hợp pháp.")
    ]
    add_drawio_file_box(s7, Inches(0.8), Inches(3.78), Inches(5.50), Inches(2.45),
                        "rtl/core/picorv32.v [RISC-V CPU Core]", code_cpu, exp_cpu, code_ratio=0.52)

    # Box 3: sections.lds (Top Right)
    code_lds = [
        "/* firmware/sections.lds - Phân bổ bộ nhớ Linker Script */",
        "MEMORY {",
        "  FLASH (rx) : ORIGIN = 0x00250000, LENGTH = 0x000B0000 /* 704KB */",
        "  RAM (rwx)  : ORIGIN = 0x00000000, LENGTH = 0x00000400 /* 1KB */",
        "}",
        ".text : { *(.text.start) *(.text*) } > FLASH",
        "_stack_top = 0x00000400; /* Khớp STACKADDR phần cứng */"
    ]
    exp_lds = [
        ("FLASH ORIGIN = 0x0025_0000", "Định vị toàn bộ mã máy C (.text) vào Flash XIP khớp Vector Reset RTL."),
        ("RAM ORIGIN = 0x0000_0000", "Dành 1KB SRAM cho biến toàn cục (.data, .bss) và Stack frame."),
        ("_stack_top = 0x0400", "Xuất nhãn đỉnh ngăn xếp cho assembly start.s nạp vào thanh ghi sp.")
    ]
    add_drawio_file_box(s7, Inches(7.033), Inches(1.22), Inches(5.50), Inches(2.45),
                        "firmware/sections.lds [Linker Script]", code_lds, exp_lds, code_ratio=0.55)

    # Box 4: start.s (Bottom Right)
    code_start = [
        "# firmware/start.s - Assembly Bootstrap khởi tạo CPU",
        ".global _start",
        "_start: lui  sp, %hi(_stack_top)       # Nạp 20-bit cao (0x00000400)",
        "        addi sp, sp, %lo(_stack_top)   # Nạp 12-bit thấp",
        "        call main                      # Chuyển quyền điều khiển sang C",
        "        ebreak; j _start               # Trap bảo vệ nếu main return"
    ]
    exp_start = [
        ("lui & addi sp, _stack_top", "Nạp mốc _stack_top (0x0400) vào con trỏ ngăn xếp sp của CPU."),
        ("call main", "Chuyển giao quyền điều khiển từ assembly sang hàm main() viết bằng C."),
        ("Vòng lặp an toàn", "Bảo vệ CPU không bị treo bus nếu chương trình chính kết thúc.")
    ]
    add_drawio_file_box(s7, Inches(7.033), Inches(3.78), Inches(5.50), Inches(2.45),
                        "firmware/start.s [Assembly Bootstrap]", code_start, exp_start, code_ratio=0.55)

    # Connectors Slide 1 (Draw.io B&W Arrows)
    add_drawio_connector(s7, Inches(6.38), Inches(1.80), Inches(0.57), Inches(0.14), direction="right", label="Boot Vector")
    add_drawio_connector(s7, Inches(6.38), Inches(4.30), Inches(0.57), Inches(0.14), direction="right", label="Stack Top")

    add_bottom_summary_box(s7, "VECTOR RESET FLASH XIP & CON TRỎ NGĂN XẾP SRAM",
                           "• Khớp Vector Khởi Động: PROGADDR_RESET (RTL) = ORIGIN FLASH (Linker) = 0x0025_0000 -> CPU nhả reset là kéo lệnh opcode thực thi trực tiếp từ Flash mà không cần RAM.\n• Khớp Ngăn Xếp: STACKADDR (RTL) = _stack_top (Linker) = sp (start.s) = 0x0000_0400 -> Biến cục bộ C truy xuất trên 1KB SRAM nội bộ cực nhanh trong 1 chu kỳ clock.")

    # =========================================================================
    # SLIDE 2 (Slide 08): BỘ ĐIỀU KHIỂN SPI FLASH XIP & CODE GHI FLASH TỪ RAM
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Phần 5: Cấu Hình Phần Cứng & Khởi Động Firmware (Slide 2/3)",
               "Co-Design SPI Flash: Bộ Điều Khiển spimemio.v & Cơ Chế Thực Thi Mã RAM", 8)

    # Box 1: spimemio.v (Left Column, Large)
    code_spimem = [
        "// rtl/core/spimemio.v - Bộ điều khiển Flash SPI QSPI / XIP",
        "module spimemio (",
        "  input valid; output ready; input [23:0] addr; output reg [31:0] rdata;",
        "  output flash_csb, flash_clk; inout [3:0] flash_io;",
        "  input [3:0] cfgreg_we; input [31:0] cfgreg_di; output [31:0] cfgreg_do",
        ");",
        "assign ready = valid && (addr == rd_addr) && rd_valid; // XIP Ready",
        "assign cfgreg_do[31]=config_en; assign cfgreg_do[5]=flash_csb; // Bit-bang"
    ]
    exp_spimem = [
        ("Giao tiếp bus XIP", "Đọc trong suốt: CPU phát địa chỉ addr[23:0], khối tự sinh xung SPI trả opcode về rdata."),
        ("4 chân vật lý SPI", "flash_csb, flash_clk, flash_io[3:0] kết nối trực tiếp chip Flash SPI ngoài bo mạch."),
        ("Kênh Bit-bang MMIO", "cfgreg_we/di/do cho phép phần mềm C điều khiển trực tiếp từng chân SPI để ghi xóa sector.")
    ]
    add_drawio_file_box(s8, Inches(0.8), Inches(1.22), Inches(5.50), Inches(5.00),
                        "rtl/core/spimemio.v [SPI Flash Controller]", code_spimem, exp_spimem, code_ratio=0.50)

    # Box 2: soc_interconnect.v Flash Decoding (Top Right)
    code_ic_flash = [
        "// rtl/core/soc_interconnect.v - Phân luồng Flash XIP & Cfg MMIO",
        "assign sel_spimem = cpu_mem_valid && (cpu_mem_addr >= 32'h0010_0000 &&",
        "                                      cpu_mem_addr <  32'h0100_0000);",
        "assign sel_spicfg = cpu_mem_valid && (cpu_mem_addr == 32'h0200_0000);",
        "assign cpu_mem_rdata = sel_spimem ? spimem_rdata :",
        "                       sel_spicfg ? spimem_cfg_do : ..."
    ]
    exp_ic_flash = [
        ("sel_spimem (XIP)", "Kích hoạt spimemio chế độ đọc lệnh trong suốt từ Flash."),
        ("sel_spicfg (MMIO)", "Kích hoạt thanh ghi cấu hình bit-bang tại địa chỉ 0x0200_0000.")
    ]
    add_drawio_file_box(s8, Inches(7.033), Inches(1.22), Inches(5.50), Inches(2.40),
                        "rtl/core/soc_interconnect.v [Flash Decoding]", code_ic_flash, exp_ic_flash, code_ratio=0.55)

    # Box 3: flash.c / start.s RAM Routine (Bottom Right)
    code_c_flash = [
        "// firmware/flash.c - Routine nạp lên SRAM để xóa/ghi Flash",
        "flashio_worker: li t0, 0x02000000; sh t1, 0(t0); # Ghi Bit-bang SPI",
        "static void flashio(uint8_t *data, int len, uint8_t wrencmd) {",
        "  // Sao chép routine flashio_worker lên Stack SRAM (< 0x0400)",
        "  // Nhảy sang SRAM thực thi lệnh bit-bang xóa sector & ghi trang Flash",
        "}"
    ]
    exp_c_flash = [
        ("Xung đột Flash XIP", "Flash không thể vừa cấp lệnh đọc XIP vừa thực thi chu kỳ ghi/xóa sector."),
        ("Thực thi từ SRAM", "CPU nhảy sang SRAM (< 0x0400) phát lệnh bit-bang 0x0200_0000 mà không nghẽn bus.")
    ]
    add_drawio_file_box(s8, Inches(7.033), Inches(3.78), Inches(5.50), Inches(2.45),
                        "firmware/flash.c & start.s [RAM Write Routine]", code_c_flash, exp_c_flash, code_ratio=0.55)

    # Connectors Slide 2 (Draw.io B&W Arrows)
    add_drawio_connector(s8, Inches(6.38), Inches(1.80), Inches(0.57), Inches(0.14), direction="left", label="Opcode XIP")
    add_drawio_connector(s8, Inches(6.38), Inches(4.30), Inches(0.57), Inches(0.14), direction="left", label="Bit-bang MMIO")

    add_bottom_summary_box(s8, "GIẢI QUYẾT XUNG ĐỘT FLASH XIP: NẠP VÀO RAM ĐỂ GHI FLASH",
                           "• Chế Độ Đọc XIP (0x0010_0000): soc_interconnect kích hoạt sel_spimem để CPU kéo opcode liên tục trực tiếp từ Flash.\n• Chế Độ Ghi Flash (0x0200_0000): flash.c nạp routine flashio_worker lên SRAM (< 0x0400) rồi chuyển quyền CPU sang SRAM để bit-bang ghi Flash Whitelist (Sec 48) và Log (Sec 49) mà không gây stall CPU.")

    # =========================================================================
    # SLIDE 3 (Slide 09): KHỐI BUS INTERCONNECT & NGOẠI VI MMIO
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Phần 5: Cấu Hình Phần Cứng & Khởi Động Firmware (Slide 3/3)",
               "Co-Design Ngoại Vi: Khối Interconnect Giải Mã MMIO Dual UART & GPIO", 9)

    # Box 1: soc_interconnect.v MMIO (Left Column, Large)
    code_ic_mmio = [
        "// rtl/core/soc_interconnect.v - Giải mã địa chỉ tổ hợp 0-delay",
        "assign sel_sram = cpu_mem_valid && (cpu_mem_addr < 32'h0000_0400);",
        "assign sel_rfid = cpu_mem_valid && (cpu_mem_addr[31:28] == 4'h1); // 0x1000_0000",
        "assign sel_uart = cpu_mem_valid && (cpu_mem_addr[31:28] == 4'h3); // 0x3000_0000",
        "assign sel_gpio = cpu_mem_valid && (cpu_mem_addr[31:28] == 4'h4); // 0x4000_0000",
        "assign cpu_mem_ready = sram_ready || spimem_ready || sel_spicfg ||",
        "                       rfid_ready || uart_ready   || gpio_ready;"
    ]
    exp_ic_mmio = [
        ("Giải mã tiền tố 4-bit cao", "Phân luồng bus tức thời trong 0 chu kỳ clock theo cpu_mem_addr[31:28]."),
        ("Tách bạch Bus", "SRAM (Stack), Flash (XIP), RFID UART, PC UART và GPIO LED hoàn toàn độc lập."),
        ("Bắt tay 1 Chu Kỳ", "Mọi ngoại vi MMIO phản hồi ready = 1 tức thì, không bao giờ làm stall CPU.")
    ]
    add_drawio_file_box(s9, Inches(0.8), Inches(1.22), Inches(5.50), Inches(5.00),
                        "rtl/core/soc_interconnect.v [MMIO Decoder & Mux]", code_ic_mmio, exp_ic_mmio, code_ratio=0.50)

    # Box 2: uart_mmio.v Dual UART (Top Right)
    code_uart = [
        "// rtl/uart/uart_mmio.v - Khối UART MMIO tích hợp FIFO 32B",
        "wire reg_div_sel = valid && (addr[2] == 1'b0); // 0x00: Baud Divider",
        "wire reg_dat_sel = valid && (addr[2] == 1'b1); // 0x04: Data FIFO",
        "assign reg_dat_do = fifo_empty ? 32'hFFFFFFFF : {24'd0, fifo_dout};",
        "assign ready = 1'b1; // Zero wait-state MMIO response"
    ]
    exp_uart = [
        ("Thanh ghi Baud (0x00)", "Cài đặt tốc độ truyền: 5208 = 50 MHz / 9600 bps."),
        ("FIFO RX 32 Byte", "Đọc không khóa: có dữ liệu trả byte mã thẻ, rỗng trả về 0xFFFFFFFF.")
    ]
    add_drawio_file_box(s9, Inches(7.033), Inches(1.22), Inches(5.50), Inches(2.40),
                        "rtl/uart/uart_mmio.v [Dual UART FIFO 32B]", code_uart, exp_uart, code_ratio=0.55)

    # Box 3: soc_regs.h & access_control.c (Bottom Right)
    code_c_driver = [
        "// firmware/common/soc_regs.h & access_control.c - Driver MMIO",
        "#define REG_RFID_UART_DAT (*(volatile uint32_t*)0x10000004)",
        "#define REG_PC_UART_DAT   (*(volatile uint32_t*)0x30000004)",
        "#define REG_GPIO_LEDS     (*(volatile uint32_t*)0x40000000)",
        "uint32_t d = REG_RFID_UART_DAT;",
        "if (d != 0xFFFFFFFF) rdm6300_push((uint8_t)d); // Polling không stall"
    ]
    exp_c_driver = [
        ("Con trỏ volatile uint32_t*", "Ép trình biên dịch GCC phát sinh lệnh lw/sw trực tiếp trên bus MMIO."),
        ("Vòng lặp Polling Non-blocking", "CPU kiểm tra liên tục trạng thái thẻ và lệnh PC mà không bị treo hệ thống.")
    ]
    add_drawio_file_box(s9, Inches(7.033), Inches(3.78), Inches(5.50), Inches(2.45),
                        "firmware/common/soc_regs.h [Firmware Driver]", code_c_driver, exp_c_driver, code_ratio=0.55)

    # Connectors Slide 3 (Draw.io B&W Arrows)
    add_drawio_connector(s9, Inches(6.38), Inches(1.80), Inches(0.57), Inches(0.14), direction="right", label="Ánh Xạ MMIO")
    add_drawio_connector(s9, Inches(6.38), Inches(4.30), Inches(0.57), Inches(0.14), direction="right", label="Thanh Ghi C")

    add_bottom_summary_box(s9, "BẮT TAY MMIO TRONG SUỐT VÀ GIAO TIẾP NGOẠI VI KHÔNG KHÓA",
                           "• Bắt tay 1 Chu Kỳ (Zero-wait-state): Mọi thanh ghi MMIO trả ready = 1 tức thì, CPU PicoRV32 luôn vận hành ở hiệu năng cao nhất.\n• FIFO Đệm Phần Cứng 32 Byte: Đầu đọc RFID phát chuỗi 14 byte ASCII, phần cứng tự động nạp vào FIFO; firmware chỉ cần đọc thanh ghi 0x1000_0004 để xử lý mã thẻ.")

    out_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "test_3slides.pptx")
    prs.save(out_file)
    print(f"Generated test deck: {out_file}")

if __name__ == "__main__":
    build_test()
